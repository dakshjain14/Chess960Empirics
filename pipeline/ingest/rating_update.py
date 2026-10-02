"""rating_update.py — refreshes Freestyle PGN WhiteElo/BlackElo headers from
data/Rating_lists/ CSVs (see Data_Selection.md's FIDE-enrichment section for
why this exists).

Matches each side to its roster CSV row by WhiteFideId/BlackFideId first,
falling back to a normalised name match; sets Elo to the resolved
Fide_Rating, to "0" if the row exists but has no Fide_id (no FIDE profile to
rate them by), or leaves it untouched with a printed ALERT if the player
isn't in the CSV at all.

Edits are plain-text tag substitution rather than a chess.pgn round-trip,
since python-chess's parser rejects some of these files' SAN as illegal (a
Chess960 castling-notation edge case) and silently truncates movetext on
export — text patching keeps every other byte identical.

data/games/ is read-only; output mirrors it under
data/processed/Updated_Ratings/Freestyle_rating_updated/.
"""
import csv
import glob
import os
import re

import openpyxl

from pipeline.config import DATA_GAMES_DIR, DATA_RATING_LISTS_DIR, FREESTYLE_MANIFEST_PATH, UPDATED_RATINGS_DIR

GAMES_ROOT = str(DATA_GAMES_DIR / 'Freestyle')
OUTPUT_ROOT = str(UPDATED_RATINGS_DIR / 'Freestyle_rating_updated')
MANIFEST_PATH = str(FREESTYLE_MANIFEST_PATH)
RATING_LISTS_ROOT = str(DATA_RATING_LISTS_DIR)

GAME_SPLIT_RE = re.compile(r'\n\n(?=\[Event )')
TAG_RE = re.compile(r'^\[(\w+)\s+"(.*)"\]\s*$', re.MULTILINE)


def find_player_list_file(filename):
    """Search all Rating_lists subdirectories for filename + '.csv' — does
    not rely on the manifest's Player_list_directory field, which does not
    track the file's actual current location."""
    matches = glob.glob(os.path.join(RATING_LISTS_ROOT, '*', filename + '.csv'))
    return matches[0] if matches else None


def load_manifest_pgn_map():
    """Return (pgn_map, known_basenames): pgn_map is basename ->
    Player_list_file_name for rows with both; known_basenames covers every
    basename with a Filepath, so the caller can tell "not in the manifest"
    apart from "in the manifest, no player-list CSV linked" (some events
    already carry complete Elo/FideId in the source PGN)."""
    wb = openpyxl.load_workbook(MANIFEST_PATH, data_only=True)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {h: i for i, h in enumerate(headers)}
    mapping = {}
    known_basenames = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        fp = row[idx['Filepath']]
        if not fp:
            continue
        basename = os.path.basename(fp)
        known_basenames.add(basename)
        fn = row[idx['Player_list_file_name']]
        if not fn:
            continue
        if basename in mapping and mapping[basename] != fn:
            print(f"ALERT: '{basename}' maps to multiple player lists in the "
                  f"manifest: {mapping[basename]!r} and {fn!r} - keeping first")
            continue
        mapping[basename] = fn
    return mapping, known_basenames


def find_header_row_flexible(rows):
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            if cell.strip() in ('Player', 'Name'):
                return i, j
    raise ValueError("no header row with a Player/Name column found")


def normalize_name(name):
    """Token-set normalisation so 'Caruana, Fabiano', 'Fabiano Caruana', and
    'Caruana Fabiano' (no comma) all compare equal."""
    tokens = name.replace(',', ' ').split()
    return frozenset(t.lower() for t in tokens)


def load_rating_lookup(player_list_filename):
    """Return (by_fide_id, by_name, no_fide_id_names), or (None, None, None)
    if the file can't be found or lacks the needed columns. by_fide_id/
    by_name map to Fide_Rating string; no_fide_id_names is the set of
    normalised names whose Fide_id is blank/'NOT FOUND' - in the roster but
    with no FIDE profile to rate them by."""
    path = find_player_list_file(player_list_filename)
    if not path:
        return None, None, None
    with open(path, encoding='utf-8-sig', newline='') as f:
        rows = list(csv.reader(f))
    header_idx, player_col = find_header_row_flexible(rows)
    header_lower = [h.strip().lower() for h in rows[header_idx]]
    if 'fide_id' not in header_lower or 'fide_rating' not in header_lower:
        return None, None, None
    fide_col = header_lower.index('fide_id')
    rating_col = header_lower.index('fide_rating')

    by_fide_id = {}
    by_name = {}
    no_fide_id_names = set()
    for row in rows[header_idx + 1:]:
        if not any(c.strip() for c in row):
            continue
        name = row[player_col].strip() if player_col < len(row) else ''
        fid = row[fide_col].strip() if fide_col < len(row) else ''
        rating = row[rating_col].strip() if rating_col < len(row) else ''
        has_fide_id = bool(fid) and fid != 'NOT FOUND'
        if not has_fide_id and name:
            no_fide_id_names.add(normalize_name(name))
            continue
        if not rating:
            continue
        by_fide_id[fid] = rating
        if name:
            by_name[normalize_name(name)] = rating
    return by_fide_id, by_name, no_fide_id_names


def lookup_rating(by_fide_id, by_name, no_fide_id_names, fide_id, name):
    """Returns (rating_or_None, reason) where reason is 'matched',
    'no_fide_id' (row exists but has no Fide_id - caller should write '0'),
    or 'unmatched' (no row at all for this player)."""
    if fide_id and fide_id in by_fide_id:
        return by_fide_id[fide_id], 'matched'
    rating = by_name.get(normalize_name(name))
    if rating:
        return rating, 'matched'
    if normalize_name(name) in no_fide_id_names:
        return None, 'no_fide_id'
    return None, 'unmatched'


def parse_tags(block):
    """First-occurrence tag -> value map for one game's raw text block."""
    tags = {}
    for m in TAG_RE.finditer(block):
        tags.setdefault(m.group(1), m.group(2))
    return tags


def set_tag(block, tag, new_value, anchor_tags):
    """Set '[tag "new_value"]' in block: replaces the value of the first
    existing '[tag "..."]' line if present (leaving everything else -
    including any duplicate later tag of the same name, movetext, comments -
    byte-identical); otherwise inserts a new tag line immediately after the
    first anchor tag in anchor_tags found present, since some games in this
    corpus omit WhiteElo/BlackElo entirely rather than using a placeholder."""
    pattern = re.compile(r'^\[' + re.escape(tag) + r'\s+"(.*)"\]\s*$', re.MULTILINE)
    if pattern.search(block):
        return pattern.sub(lambda m: m.group(0).replace(m.group(1), new_value, 1), block, count=1)
    for anchor in anchor_tags:
        anchor_pattern = re.compile(r'^(\[' + re.escape(anchor) + r'\s+".*"\])\s*$', re.MULTILINE)
        m = anchor_pattern.search(block)
        if m:
            insert_at = m.end()
            new_line = f'\n[{tag} "{new_value}"]'
            return block[:insert_at] + new_line + block[insert_at:]
    return block


def process_pgn(src_path, dst_path, by_fide_id, by_name, no_fide_id_names, basename):
    with open(src_path, encoding='utf-8-sig') as f:
        content = f.read()
    blocks = [b for b in GAME_SPLIT_RE.split(content) if b.strip()]

    n_games = 0
    n_updated = 0
    n_zeroed = 0
    n_unmatched = 0
    out_blocks = []
    for block in blocks:
        n_games += 1
        tags = parse_tags(block)
        new_block = block
        for color, elo_key, fide_key in (
            ('White', 'WhiteElo', 'WhiteFideId'),
            ('Black', 'BlackElo', 'BlackFideId'),
        ):
            name = tags.get(color, '')
            fide_id = tags.get(fide_key, '').strip()
            rating, reason = lookup_rating(by_fide_id, by_name, no_fide_id_names, fide_id, name)
            anchor_tags = (fide_key, color)  # insert after FideId if present, else after the name tag
            if reason == 'matched':
                new_block = set_tag(new_block, elo_key, rating, anchor_tags)
                n_updated += 1
            elif reason == 'no_fide_id':
                new_block = set_tag(new_block, elo_key, '0', anchor_tags)
                n_zeroed += 1
            else:
                n_unmatched += 1
                print(f"    ALERT: {basename} game {n_games}: could not resolve "
                      f"Fide rating for {color} '{name}' (FideId={fide_id or 'none'}) "
                      f"- Elo left unchanged")
        out_blocks.append(new_block)

    os.makedirs(os.path.dirname(dst_path), exist_ok=True)
    with open(dst_path, 'w', encoding='utf-8') as f:
        f.write('\n\n'.join(out_blocks))
        if content.endswith('\n') and not out_blocks[-1].endswith('\n'):
            f.write('\n')

    return n_games, n_updated, n_zeroed, n_unmatched


def main():
    pgn_map, known_basenames = load_manifest_pgn_map()
    pgn_files = sorted(glob.glob(os.path.join(GAMES_ROOT, '**', '*.pgn'), recursive=True))
    print(f"found {len(pgn_files)} Freestyle PGN files on disk")

    rating_cache = {}  # player_list_filename -> (by_fide_id, by_name, no_fide_id_names)

    for src_path in pgn_files:
        basename = os.path.basename(src_path)
        if basename not in known_basenames:
            print(f"ALERT: '{basename}' not referenced in freestyle_manifest.xlsx - skipping, no output written")
            continue

        player_list_fn = pgn_map.get(basename)
        if not player_list_fn:
            # In the manifest, but no player-list CSV linked - nothing to
            # backfill from. Copy the file through unchanged rather than
            # skipping it, so it still reaches data/processed/ for the
            # time-annotation step; these events typically already carry
            # complete WhiteElo/BlackElo/WhiteFideId/BlackFideId in the
            # source PGN.
            rel = os.path.relpath(src_path, GAMES_ROOT)
            dst_path = os.path.join(OUTPUT_ROOT, rel)
            os.makedirs(os.path.dirname(dst_path), exist_ok=True)
            with open(src_path, encoding='utf-8-sig') as f:
                content = f.read()
            with open(dst_path, 'w', encoding='utf-8') as f:
                f.write(content)
            print(f"{basename}: no player list linked - copied through unchanged -> {dst_path}")
            continue

        if player_list_fn not in rating_cache:
            rating_cache[player_list_fn] = load_rating_lookup(player_list_fn)
        by_fide_id, by_name, no_fide_id_names = rating_cache[player_list_fn]
        if by_fide_id is None:
            print(f"ALERT: player list '{player_list_fn}' for '{basename}' not found on disk "
                  f"or missing Fide_id/Fide_Rating columns - skipping, no output written")
            continue

        rel = os.path.relpath(src_path, GAMES_ROOT)
        dst_path = os.path.join(OUTPUT_ROOT, rel)
        n_games, n_updated, n_zeroed, n_unmatched = process_pgn(
            src_path, dst_path, by_fide_id, by_name, no_fide_id_names, basename)
        msg = f"{basename}: {n_games} games, {n_updated} sides updated"
        if n_zeroed:
            msg += f", {n_zeroed} sides set to 0 (in player list, no Fide_id)"
        msg += f" -> {dst_path}"
        if n_unmatched:
            msg += f" - {n_unmatched} sides left unmatched (see ALERTs above)"
        print(msg)

    print("done")


if __name__ == '__main__':
    main()
