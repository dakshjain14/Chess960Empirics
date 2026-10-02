"""twic_event_stats.py — recomputes twic_event_classification.xlsx's
per-event stats directly from TWIC's PGN archive (path/to/twic/*.pgn, not
tracked in this repo). See Data_Selection.md's "Source" and "Median-max
rating rule" sections.

Computes Total_games, Chess960_games, Classification (best-effort — TWIC
PGNs carry no TimeControl tag, so this isn't authoritative and needs manual
review), Median_max, Games_2100_2399, and the rating/player-count columns
from Data_Selection.md's floor rules. Leaves every MANUAL_COLUMNS cell
(Study_status, thresholds, availability, rationale, ...) untouched on
existing rows and blank on new ones — those need judgment not recoverable
from the PGN alone.

Run:
    PYTHONPATH=. .venv-pipeline/bin/python pipeline/ingest/twic_event_stats.py \\
        --twic-dir path/to/twic/ --out data/games/twic_event_classification.xlsx
"""
from __future__ import annotations

import argparse
import glob
import re
import statistics
from collections import defaultdict
from pathlib import Path

import openpyxl

from pipeline.config import DATA_GAMES_DIR

DEFAULT_OUT = DATA_GAMES_DIR / "twic_event_classification.xlsx"

TAG_RE = re.compile(r'^\[(\w+)\s+"(.*)"\]\s*$', re.MULTILINE)
GAME_SPLIT_RE = re.compile(r"\n\n(?=\[Event )")

MANUAL_COLUMNS = (
    "Study_status",
    "Game Threshold rule (30 for classical, 25 for Rapid)",
    "Game Availability",
    "In_manifest",
    "Rationale",
    "Comment",
)

COMPUTED_COLUMNS = (
    "Total_games",
    "Chess960_games",
    "Classification",
    "Total_games_verified",
    "Median_max",
    "Games_2100_2399",
    "games_both_2000plus",
    "median_rating",
    "% games of 2000+ players",
    "total_players",
    "players_2000plus",
    "players_2400plus",
    "% player above 2400",
)

ALL_COLUMNS = (
    "Event", "Total_games", "Chess960_games", "Classification", "Study_status",
    "Total_games_verified", "Game Threshold rule (30 for classical, 25 for Rapid)",
    "Game Availability", "Median_max", "In_manifest", "Games_2100_2399",
    "Rationale", "Comment", "games_both_2000plus", "median_rating",
    "% games of 2000+ players", "total_players", "players_2000plus",
    "players_2400plus", "% player above 2400",
)


def _to_int(x):
    try:
        return int(x)
    except (TypeError, ValueError):
        return None


def parse_pgn_games(path):
    """Yield each game's tag dict (Event, White, Black, WhiteElo, BlackElo,
    Variant, ...) from a TWIC PGN file."""
    with open(path, encoding="utf-8", errors="replace") as f:
        content = f.read()
    for block in GAME_SPLIT_RE.split(content):
        if not block.strip():
            continue
        tags = {}
        for m in TAG_RE.finditer(block):
            tags.setdefault(m.group(1), m.group(2))
        if "Event" in tags:
            yield tags


def group_games_by_event(twic_dir):
    """Read every *.pgn directly under twic_dir, return {event_name: [tags, ...]}."""
    by_event = defaultdict(list)
    for path in sorted(glob.glob(str(Path(twic_dir) / "*.pgn"))):
        for tags in parse_pgn_games(path):
            by_event[tags["Event"].strip()].append(tags)
    return by_event


def classify(event_name, games):
    """Best-effort Classification — see module docstring. Not authoritative."""
    if games and all(g.get("Variant", "").strip().lower() == "chess960" for g in games):
        return "Chess_960"
    name_lower = event_name.lower()
    if "rapid" in name_lower:
        return "Rapid"
    if "blitz" in name_lower:
        return "Blitz"
    return "Classical"


def compute_event_stats(event_name, games):
    """Compute every COMPUTED_COLUMNS value for one event (pure, no I/O)."""
    total_games = len(games)
    chess960_games = sum(1 for g in games if g.get("Variant", "").strip().lower() == "chess960")

    maxes = []
    all_ratings = []
    both_2000plus = 0
    in_band = 0
    players = {}  # name -> last-seen rating
    for g in games:
        we = _to_int(g.get("WhiteElo"))
        be = _to_int(g.get("BlackElo"))
        if we is not None and be is not None:
            mx = max(we, be)
            maxes.append(mx)
            if we >= 2000 and be >= 2000:
                both_2000plus += 1
            if 2100 <= mx <= 2399:
                in_band += 1
        if we is not None:
            all_ratings.append(we)
        white = g.get("White")
        if white and we is not None:
            players[white] = we
        black = g.get("Black")
        if black and be is not None:
            players[black] = be
        if be is not None:
            all_ratings.append(be)

    total_players = len(players)
    players_2000plus = sum(1 for r in players.values() if r >= 2000)
    players_2400plus = sum(1 for r in players.values() if r >= 2400)

    return {
        "Total_games": total_games,
        "Chess960_games": chess960_games,
        "Classification": classify(event_name, games),
        "Total_games_verified": total_games,
        "Median_max": statistics.median(maxes) if maxes else None,
        "Games_2100_2399": in_band,
        "games_both_2000plus": both_2000plus,
        "median_rating": statistics.median(all_ratings) if all_ratings else None,
        "% games of 2000+ players": (both_2000plus / total_games) if total_games else None,
        "total_players": total_players,
        "players_2000plus": players_2000plus,
        "players_2400plus": players_2400plus,
        "% player above 2400": (players_2400plus / total_players) if total_players else None,
    }


def merge_into_workbook(out_path, stats_by_event):
    """Update COMPUTED_COLUMNS for existing rows (MANUAL_COLUMNS untouched)
    and append new rows for events not yet in the sheet."""
    out_path = Path(out_path)
    if out_path.exists():
        wb = openpyxl.load_workbook(out_path)
        ws = wb.active
        header_row = next(ws.iter_rows(min_row=1, max_row=1))
        headers = [c.value for c in header_row]
        col_idx = {h: i + 1 for i, h in enumerate(headers) if h}
        for name in ALL_COLUMNS:
            if name not in col_idx:
                raise RuntimeError(f"{out_path} is missing expected column {name!r}")

        event_col = col_idx["Event"]
        existing_rows = {}
        for row in ws.iter_rows(min_row=2):
            ev = row[event_col - 1].value
            if ev:
                existing_rows[str(ev).strip()] = row[0].row

        updated, added = 0, 0
        for event, stats in stats_by_event.items():
            if event in existing_rows:
                r = existing_rows[event]
                for col_name in COMPUTED_COLUMNS:
                    ws.cell(row=r, column=col_idx[col_name], value=stats[col_name])
                updated += 1
            else:
                r = ws.max_row + 1
                ws.cell(row=r, column=event_col, value=event)
                for col_name in COMPUTED_COLUMNS:
                    ws.cell(row=r, column=col_idx[col_name], value=stats[col_name])
                added += 1
        wb.save(out_path)
        return updated, added

    wb = openpyxl.Workbook()
    ws = wb.active
    col_idx = {name: i + 1 for i, name in enumerate(ALL_COLUMNS)}
    for name, i in col_idx.items():
        ws.cell(row=1, column=i, value=name)
    r = 2
    for event, stats in stats_by_event.items():
        ws.cell(row=r, column=col_idx["Event"], value=event)
        for col_name in COMPUTED_COLUMNS:
            ws.cell(row=r, column=col_idx[col_name], value=stats[col_name])
        r += 1
    wb.save(out_path)
    return 0, len(stats_by_event)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--twic-dir", required=True, help="Directory of TWIC *.pgn bulletin files")
    parser.add_argument("--out", default=str(DEFAULT_OUT), help="Target xlsx (merged in place if it exists)")
    args = parser.parse_args()

    by_event = group_games_by_event(args.twic_dir)
    print(f"[twic_event_stats] parsed {len(by_event)} events from {args.twic_dir}")

    stats_by_event = {event: compute_event_stats(event, games) for event, games in by_event.items()}

    updated, added = merge_into_workbook(args.out, stats_by_event)
    print(f"[twic_event_stats] {args.out}: {updated} existing rows updated, {added} new rows added")
    print("[twic_event_stats] Classification is a best-effort heuristic (no TimeControl "
          "tag in TWIC PGNs) — review it manually, along with every column in "
          "MANUAL_COLUMNS, for any newly added rows.")


if __name__ == "__main__":
    main()
