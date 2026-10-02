"""Tests for pipeline/ingest/twic_event_stats.py.

Run:  .venv-pipeline/bin/python -m pytest pipeline/tests/test_twic_event_stats.py -q
"""
from __future__ import annotations

import openpyxl
import pytest

from pipeline.ingest.twic_event_stats import (
    ALL_COLUMNS,
    classify,
    compute_event_stats,
    merge_into_workbook,
    parse_pgn_games,
)


def _game(white, we, black, be, variant=None):
    tags = {"Event": "Test Event", "White": white, "WhiteElo": str(we),
            "Black": black, "BlackElo": str(be)}
    if variant:
        tags["Variant"] = variant
    return tags


def test_compute_event_stats_basic_counts():
    games = [
        _game("Alice", 2200, "Bob", 2350),
        _game("Bob", 2350, "Carol", 1900),
        _game("Carol", 1900, "Alice", 2200),
    ]
    stats = compute_event_stats("Test Event", games)
    assert stats["Total_games"] == 3
    assert stats["Total_games_verified"] == 3
    assert stats["Chess960_games"] == 0
    assert stats["Median_max"] == 2350
    assert stats["total_players"] == 3
    assert stats["players_2000plus"] == 2
    assert stats["players_2400plus"] == 0


def test_compute_event_stats_median_and_band():
    # per-game max(White,Black): 2350, 2350, 2200 -> sorted [2200,2350,2350] -> median 2350
    games = [
        _game("Alice", 2200, "Bob", 2350),
        _game("Bob", 2350, "Carol", 1900),
        _game("Carol", 1900, "Alice", 2200),
    ]
    stats = compute_event_stats("Test Event", games)
    assert stats["Median_max"] == 2350
    # 2100-2399 band: all three per-game maxes (2350, 2350, 2200) fall in [2100,2399]
    assert stats["Games_2100_2399"] == 3
    assert stats["games_both_2000plus"] == 1  # only Alice/Bob game has both >=2000


def test_compute_event_stats_empty_input_safe():
    stats = compute_event_stats("Empty Event", [])
    assert stats["Total_games"] == 0
    assert stats["Median_max"] is None
    assert stats["median_rating"] is None
    assert stats["% games of 2000+ players"] is None
    assert stats["% player above 2400"] is None


def test_classify_chess960():
    games = [_game("A", 2200, "B", 2200, variant="Chess960")]
    assert classify("Some 960 Open", games) == "Chess_960"


def test_classify_rapid_keyword():
    games = [_game("A", 2200, "B", 2200)]
    assert classify("City Rapid Championship", games) == "Rapid"


def test_classify_blitz_keyword():
    games = [_game("A", 2200, "B", 2200)]
    assert classify("Friday Night Blitz", games) == "Blitz"


def test_classify_default_classical():
    games = [_game("A", 2200, "B", 2200)]
    assert classify("City Open 2025", games) == "Classical"


def test_parse_pgn_games_splits_on_event_tag(tmp_path):
    pgn_text = (
        '[Event "Game One"]\n[White "A"]\n[WhiteElo "2200"]\n'
        '[Black "B"]\n[BlackElo "2300"]\n\n1. e4 e5 *\n\n'
        '[Event "Game Two"]\n[White "C"]\n[WhiteElo "2100"]\n'
        '[Black "D"]\n[BlackElo "2400"]\n\n1. d4 d5 *\n'
    )
    path = tmp_path / "sample.pgn"
    path.write_text(pgn_text)
    games = list(parse_pgn_games(path))
    assert len(games) == 2
    assert games[0]["Event"] == "Game One"
    assert games[1]["White"] == "C"


def _write_workbook(path, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    for i, name in enumerate(ALL_COLUMNS, start=1):
        ws.cell(row=1, column=i, value=name)
    for r, row in enumerate(rows, start=2):
        for i, name in enumerate(ALL_COLUMNS, start=1):
            ws.cell(row=r, column=i, value=row.get(name))
    wb.save(path)
    return path


def test_merge_into_workbook_preserves_manual_columns(tmp_path):
    out_path = tmp_path / "twic.xlsx"
    _write_workbook(out_path, [{
        "Event": "Existing Event",
        "Total_games": 10,
        "Study_status": "Qualified",
        "Rationale": "Included",
        "Comment": "hand-reviewed",
        "In_manifest": "Yes",
    }])

    new_stats = {"Existing Event": compute_event_stats("Existing Event", [
        _game("Alice", 2500, "Bob", 2600),
    ])}
    updated, added = merge_into_workbook(out_path, new_stats)
    assert updated == 1
    assert added == 0

    wb = openpyxl.load_workbook(out_path)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {h: i for i, h in enumerate(headers)}
    row = next(ws.iter_rows(min_row=2, max_row=2, values_only=True))
    assert row[idx["Total_games"]] == 1  # computed column overwritten
    assert row[idx["Study_status"]] == "Qualified"  # manual column untouched
    assert row[idx["Rationale"]] == "Included"
    assert row[idx["Comment"]] == "hand-reviewed"
    assert row[idx["In_manifest"]] == "Yes"


def test_merge_into_workbook_appends_new_event_with_blank_manual_columns(tmp_path):
    out_path = tmp_path / "twic.xlsx"
    _write_workbook(out_path, [{"Event": "Existing Event", "Total_games": 10, "Study_status": "Qualified"}])

    new_stats = {"Brand New Event": compute_event_stats("Brand New Event", [
        _game("Alice", 2500, "Bob", 2600),
    ])}
    updated, added = merge_into_workbook(out_path, new_stats)
    assert updated == 0
    assert added == 1

    wb = openpyxl.load_workbook(out_path)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {h: i for i, h in enumerate(headers)}
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    new_row = next(r for r in rows if r[idx["Event"]] == "Brand New Event")
    assert new_row[idx["Total_games"]] == 1
    assert new_row[idx["Study_status"]] is None


def test_merge_into_workbook_creates_file_when_absent(tmp_path):
    out_path = tmp_path / "does_not_exist.xlsx"
    new_stats = {"Only Event": compute_event_stats("Only Event", [
        _game("Alice", 2500, "Bob", 2600),
    ])}
    updated, added = merge_into_workbook(out_path, new_stats)
    assert updated == 0
    assert added == 1
    assert out_path.exists()
